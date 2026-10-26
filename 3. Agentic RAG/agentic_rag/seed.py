"""Deterministic demo data: a SQLite DB plus QC reports, supplier specs (docx), manuals, and tickets (xlsx + txt).

Planted stories the agent should find:
  1. B-2025-07-A (CF-1200, Pune, VoltCap): capacitor lot CAP-7781 out of tolerance, released under deviation DV-221,
     then a spike of 'capacitor failure' / 'noise' returns in September 2025.
  2. B-2025-08-E (TF-400, Pune, Zenith): noise/buzzing complaints.
Everything else is low-rate background noise.
"""
import random
import sqlite3
from datetime import date, timedelta
from pathlib import Path

from .config import settings

MODELS = [  # (model, pune_letter, hanoi_letter, base_units)
    ("CF-1200", "A", "B", 4000),
    ("CF-900", "C", "D", 3000),
    ("TF-400", "E", "F", 5000),
    ("EF-150", "G", "H", 6000),
]
SUPPLIER = {("CF", "Pune"): "VoltCap", ("CF", "Hanoi"): "Kapex", ("TF", "Pune"): "Zenith", ("TF", "Hanoi"): "Kapex",
            ("EF", "Pune"): "Zenith", ("EF", "Hanoi"): "Kapex"}
REASONS = ["capacitor failure", "noise", "wobble", "remote issue", "blade damage", "switch fault"]
REGIONS = [("EU", 0.40), ("US", 0.35), ("APAC", 0.25)]
BAD_BATCH, NOISE_BATCH = "B-2025-07-A", "B-2025-08-E"

SCHEMA_SQL = """
CREATE TABLE production_batches(batch_id TEXT PRIMARY KEY, model TEXT, factory TEXT, produced_date TEXT,
                                units_produced INT, capacitor_supplier TEXT);
CREATE TABLE sales(sale_id INTEGER PRIMARY KEY, model TEXT, batch_id TEXT, region TEXT, sale_date TEXT, units_sold INT);
CREATE TABLE returns(return_id INTEGER PRIMARY KEY, model TEXT, batch_id TEXT, return_date TEXT, reason TEXT, units_returned INT);
CREATE TABLE qc_inspections(inspection_id INTEGER PRIMARY KEY, batch_id TEXT, inspection_date TEXT, units_inspected INT,
                            units_failed INT, defect_type TEXT);
"""


def _build_data():
    rng = random.Random(42)
    batches, sales, returns, qc = [], [], [], []
    rid = sid = qid = 0
    for model, pune, hanoi, base in MODELS:
        for month in range(1, 10):
            for factory, letter in (("Pune", pune), ("Hanoi", hanoi)):
                bid = f"B-2025-{month:02d}-{letter}"
                produced = date(2025, month, 8 if factory == "Pune" else 15)
                units = base + rng.randint(-300, 300)
                supplier = SUPPLIER[(model[:2], factory)]
                batches.append((bid, model, factory, produced.isoformat(), units, supplier))

                sold_total = int(units * 0.92)
                for region, share in REGIONS:
                    sid += 1
                    sales.append((sid, model, bid, region, (produced + timedelta(days=35)).isoformat(), int(sold_total * share)))

                for k, days in enumerate((60, 90)):  # background returns, ~0.7% of sold
                    for reason in rng.sample(REASONS, 2):
                        n = max(1, int(sold_total * 0.0018 * rng.uniform(0.5, 1.5)))
                        rid += 1
                        returns.append((rid, model, bid, (produced + timedelta(days=days + k * 3)).isoformat(), reason, n))

                rate = 0.14 if bid == BAD_BATCH else (0.03 if factory == "Pune" else 0.02) * rng.uniform(0.7, 1.3)
                qid += 1
                qc.append((qid, bid, (produced + timedelta(days=3)).isoformat(), 200, round(200 * rate),
                           "capacitor out of tolerance" if bid == BAD_BATCH else rng.choice(["blade balance", "housing", "wiring", "label"])))
    # Planted anomalies
    rid += 1; returns.append((rid, "CF-1200", BAD_BATCH, "2025-09-02", "capacitor failure", 310))
    rid += 1; returns.append((rid, "CF-1200", BAD_BATCH, "2025-09-20", "noise", 95))
    rid += 1; returns.append((rid, "TF-400", NOISE_BATCH, "2025-09-28", "noise", 140))
    return batches, sales, returns, qc


def build_database() -> None:
    settings.db_path.parent.mkdir(parents=True, exist_ok=True)
    if settings.db_path.exists():
        settings.db_path.unlink()
    batches, sales, returns, qc = _build_data()
    con = sqlite3.connect(settings.db_path)
    con.executescript(SCHEMA_SQL)
    con.executemany("INSERT INTO production_batches VALUES (?,?,?,?,?,?)", batches)
    con.executemany("INSERT INTO sales VALUES (?,?,?,?,?,?)", sales)
    con.executemany("INSERT INTO returns VALUES (?,?,?,?,?,?)", returns)
    con.executemany("INSERT INTO qc_inspections VALUES (?,?,?,?,?,?)", qc)
    con.commit()
    con.close()


def _write(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


def build_documents() -> None:
    from docx import Document
    from openpyxl import Workbook

    docs = settings.docs_dir
    batches, _, _, qc = _build_data()
    qc_by_batch = {q[1]: q for q in qc}

    # QC reports (markdown), one per batch
    for bid, model, factory, produced, units, supplier in batches:
        _, _, insp_date, inspected, failed, defect = qc_by_batch[bid]
        rate = 100 * failed / inspected
        lot = {"VoltCap": "CAP-7781" if bid == BAD_BATCH else "CAP-7702", "Kapex": "KX-300", "Zenith": "ZN-110"}[supplier]
        if bid == BAD_BATCH:
            findings = (f"Capacitor lot {lot} from {supplier} measured outside the +/-5% capacitance tolerance on 14% of sampled units. "
                        "Root cause suspected: revised dielectric film from a new sub-vendor.")
            disposition = "Released under deviation DV-221 despite the finding, due to shipment deadlines."
        else:
            findings = f"{failed} of {inspected} sampled units had minor cosmetic or assembly issues ({defect}). Capacitors from {supplier}, lot {lot}, within tolerance."
            disposition = "Released."
        _write(docs / "qc_reports" / f"{bid}.md", f"""# QC Report - Batch {bid}

Model: {model} | Factory: {factory} | Capacitor supplier: {supplier} | Inspection date: {insp_date}

## Summary
Inspected {inspected} units, {failed} failed ({rate:.1f}%).

## Findings
{findings}

## Disposition
{disposition}
""")

    # Supplier specs (docx, with a table)
    specs = {
        "VoltCap": ("VC-450", "2.5 uF +/-5%", "85C", "450 V",
                    "Change notice: lot CAP-7781 used a revised dielectric film from a new sub-vendor. The new film was not requalified before shipment."),
        "Kapex": ("KX-200", "2.5 uF +/-5%", "105C", "450 V", "No open change notices."),
        "Zenith": ("ZN-90", "2.0 uF +/-10%", "85C", "400 V", "No open change notices."),
    }
    for supplier, (part, cap, temp, volt, notice) in specs.items():
        d = Document()
        d.add_heading(f"{supplier} capacitor specification {part}", 1)
        d.add_paragraph(f"Film capacitor {part} supplied by {supplier} for fan motors.")
        d.add_heading("Electrical parameters", 2)
        t = d.add_table(rows=1, cols=2)
        t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Parameter", "Value"
        for k, v in (("Capacitance", cap), ("Rated temperature", temp), ("Rated voltage", volt)):
            row = t.add_row().cells
            row[0].text, row[1].text = k, v
        d.add_heading("Change notices", 2)
        d.add_paragraph(notice)
        (docs / "supplier_specs").mkdir(parents=True, exist_ok=True)
        d.save(docs / "supplier_specs" / f"{supplier}_{part}.docx")

    # Manuals
    for model, *_ in MODELS:
        _write(docs / "manuals" / f"{model}_manual.md", f"""# {model} Manual

## Installation
Mount the {model} on a rated support and connect it to a grounded supply.

## Troubleshooting: humming sound at low speed
A humming sound at low speed usually indicates a degraded start capacitor. Do not keep running the fan; contact service.

## Troubleshooting: wobble
Check blade screws and balance weights. Never bend blades.
""")

    # Tickets: an Excel log + a few individual text tickets
    wb = Workbook()
    ws = wb.active
    ws.title = "tickets"
    ws.append(["ticket_id", "date", "model", "batch_id", "complaint"])
    rng = random.Random(7)
    tid = 8000
    background = [("CF-900", "B-2025-06-C", "Blade wobbles at high speed"), ("EF-150", "B-2025-06-G", "Switch feels loose"),
                  ("CF-1200", "B-2025-05-B", "Remote not pairing")]
    for _ in range(40):
        tid += 1
        m, b, c = rng.choice(background)
        ws.append([tid, f"2025-{rng.randint(6, 9):02d}-{rng.randint(1, 28):02d}", m, b, c])
    for _ in range(35):
        tid += 1
        ws.append([tid, f"2025-09-{rng.randint(1, 28):02d}", "CF-1200", BAD_BATCH,
                   rng.choice(["Fan hums loudly at low speed and then stops", "Loud humming noise, capacitor smell", "Stops after 3 weeks, humming before failure"])])
    for _ in range(25):
        tid += 1
        ws.append([tid, f"2025-09-{rng.randint(10, 28):02d}", "TF-400", NOISE_BATCH,
                   rng.choice(["Buzzing noise from motor housing at speed 3", "Rattling noise when oscillating", "Noisy at all speeds"])])
    (docs / "service_tickets").mkdir(parents=True, exist_ok=True)
    wb.save(docs / "service_tickets" / "tickets_2025.xlsx")

    _write(docs / "service_tickets" / "ticket_9901.txt",
           f"Ticket 9901 | 2025-09-03 | CF-1200 | batch {BAD_BATCH}\nCustomer in Germany reports the fan hums loudly at low speed and stopped after 3 weeks. Technician replaced the capacitor.")
    _write(docs / "service_tickets" / "ticket_9902.txt",
           f"Ticket 9902 | 2025-09-27 | TF-400 | batch {NOISE_BATCH}\nBuzzing noise from the motor housing at speed 3. Customer requests a refund.")


def seed_all() -> dict:
    build_database()
    build_documents()
    n_docs = sum(1 for p in settings.docs_dir.rglob("*") if p.is_file())
    return {"db": str(settings.db_path), "docs_dir": str(settings.docs_dir), "documents": n_docs}
