"""Document parsing: PDF, Word, PowerPoint and plain-text formats.

Every parser returns the same shape so the rest of the pipeline doesn't care
about the file type:

    {"sections": [{"loc": "Page 3", "page": 3, "text": "...", "asset": "assets/x.jpg"}],
     "images":   [{"path": Path, "loc": "Page 3", "page": 3}],     # embedded figures
     "ocr":      [{"path": Path, "loc": "Page 7", "page": 7}],     # scanned pages to OCR
     "info":     {"pages": 12}}
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from collections import Counter
from pathlib import Path

from PIL import Image

MIN_FIGURE_PX = 120          # embedded images smaller than this (icons, bullets) are ignored
MAX_FIGURES = 200
PAGE_NUMBER = re.compile(r"^\s*(page\s*)?\d{1,4}(\s*(of|/)\s*\d{1,4})?\s*$", re.I)


def _norm(line: str) -> str:
    return re.sub(r"\d+", "#", line.strip().lower())


def _save_figure(data: bytes, dest: Path, seen: set[str]) -> Path | None:
    """Keep real figures only: decent size, sane aspect ratio, not seen before (logos repeat)."""
    digest = hashlib.sha1(data).hexdigest()
    if digest in seen:
        return None
    try:
        img = Image.open(io.BytesIO(data))
        img.load()
    except Exception:
        return None
    w, h = img.size
    if min(w, h) < MIN_FIGURE_PX or max(w, h) / max(1, min(w, h)) > 8:
        return None
    seen.add(digest)
    img = img.convert("RGB")
    img.thumbnail((1280, 1280))
    img.save(dest, "JPEG", quality=86)
    return dest


def parse_pdf(path: Path, assets: Path, fid: str) -> dict:
    import pymupdf as fitz

    doc = fitz.open(path)
    pages: list[list[str]] = []
    out = {"sections": [], "images": [], "ocr": [], "info": {"pages": doc.page_count}}
    seen: set[str] = set()

    for page in doc:
        blocks = page.get_text("blocks", sort=True)
        # Each block is one visual paragraph: join its wrapped lines back together.
        pages.append([re.sub(r"\s*\n\s*", " ", b[4]).strip() for b in blocks if b[6] == 0 and b[4].strip()])

    # Running headers/footers ("Annual Report 2024 · Page 3") repeat on most pages: drop them.
    edge = Counter()
    for blocks in pages:
        for b in {_norm(x) for x in blocks[:2] + blocks[-2:]}:
            edge[b] += 1
    threshold = max(3, int(len(pages) * 0.5))
    repeated = {k for k, v in edge.items() if v >= threshold} if len(pages) >= 3 else set()

    for i, (page, blocks) in enumerate(zip(doc, pages), 1):
        kept = [b for b in blocks if _norm(b) not in repeated and not PAGE_NUMBER.match(b)]
        thumb = assets / f"{fid}-p{i}.jpg"
        pix = page.get_pixmap(matrix=fitz.Matrix(0.5, 0.5), alpha=False)
        Image.frombytes("RGB", (pix.width, pix.height), pix.samples).save(thumb, "JPEG", quality=80)
        rel = f"assets/{thumb.name}"
        text = "\n\n".join(kept)
        if len(text) < 25:
            # Little or no text layer: probably a scanned page. Render it sharp enough to OCR.
            scan = assets / f"{fid}-p{i}-scan.png"
            page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False).save(scan)
            out["ocr"].append({"path": scan, "loc": f"Page {i}", "page": i, "asset": rel})
        else:
            out["sections"].append({"loc": f"Page {i}", "page": i, "text": text, "asset": rel})

        for img in page.get_images(full=True):
            if len(out["images"]) >= MAX_FIGURES:
                break
            try:
                data = doc.extract_image(img[0])["image"]
            except Exception:
                continue
            dest = _save_figure(data, assets / f"{fid}-p{i}-fig{len(out['images'])}.jpg", seen)
            if dest:
                out["images"].append({"path": dest, "loc": f"Page {i}", "page": i})
    out["info"]["repeated_lines_removed"] = len(repeated)
    return out


def parse_docx(path: Path, assets: Path, fid: str) -> dict:
    import docx

    d = docx.Document(str(path))
    lines: list[str] = []
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style.startswith("heading") or style == "title":
            level = int(re.sub(r"\D", "", style) or 1)
            lines.append("#" * min(level, 6) + " " + t)
        else:
            lines.append(t)
    for table in d.tables:
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        if rows:
            lines.append("\n".join(rows))

    images, seen = [], set()
    for rel in d.part.rels.values():
        if "image" in rel.reltype and len(images) < MAX_FIGURES:
            dest = _save_figure(rel.target_part.blob, assets / f"{fid}-fig{len(images)}.jpg", seen)
            if dest:
                images.append({"path": dest, "loc": "", "page": None})
    return {"sections": [{"loc": "", "page": None, "text": "\n\n".join(lines), "asset": None}],
            "images": images, "ocr": [], "info": {}}


def parse_pptx(path: Path, assets: Path, fid: str) -> dict:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(path))
    sections, images, seen = [], [], set()

    def walk(shapes):
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from walk(shape.shapes)
            else:
                yield shape

    for i, slide in enumerate(prs.slides, 1):
        parts = []
        for shape in walk(slide.shapes):
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            elif getattr(shape, "has_table", False) and shape.has_table:
                parts.append("\n".join(" | ".join(c.text.strip() for c in row.cells) for row in shape.table.rows))
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE and len(images) < MAX_FIGURES:
                try:
                    dest = _save_figure(shape.image.blob, assets / f"{fid}-s{i}-fig{len(images)}.jpg", seen)
                except Exception:
                    dest = None
                if dest:
                    images.append({"path": dest, "loc": f"Slide {i}", "page": i})
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append("Speaker notes: " + slide.notes_slide.notes_text_frame.text.strip())
        if parts:
            sections.append({"loc": f"Slide {i}", "page": i, "text": "\n\n".join(parts), "asset": None})
    return {"sections": sections, "images": images, "ocr": [], "info": {"slides": len(prs.slides)}}


def parse_text(path: Path, assets: Path, fid: str) -> dict:
    raw = path.read_bytes().decode("utf-8", errors="replace")
    ext = path.suffix.lower()
    if ext in (".csv", ".tsv"):
        # Rows become "column: value" lines so each chunk is self-describing.
        reader = csv.reader(io.StringIO(raw), delimiter="\t" if ext == ".tsv" else ",")
        rows = list(reader)
        header = rows[0] if rows else []
        raw = "\n\n".join("; ".join(f"{h}: {v}" for h, v in zip(header, r) if v.strip())
                          for r in rows[1:5001])
    elif ext == ".json":
        try:
            raw = json.dumps(json.loads(raw), indent=1, ensure_ascii=False)
        except ValueError:
            pass
    return {"sections": [{"loc": "", "page": None, "text": raw, "asset": None}],
            "images": [], "ocr": [], "info": {}}


def parse_document(path: Path, assets: Path, fid: str) -> dict:
    ext = path.suffix.lower()
    parser = {".pdf": parse_pdf, ".docx": parse_docx, ".pptx": parse_pptx}.get(ext, parse_text)
    return parser(path, assets, fid)
