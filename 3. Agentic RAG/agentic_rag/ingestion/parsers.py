"""File parsers. Each returns a list of Sections (text + page/heading + optional metadata overrides)."""
import re
from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class Section:
    text: str
    page: int | None = None
    heading: str = ""
    is_table: bool = False  # tables are never split
    meta: dict = field(default_factory=dict)  # metadata taken directly from structured sources (Excel columns)


def parse_text(path: Path) -> list[Section]:
    """Markdown / plain text: split on markdown headings."""
    sections: list[Section] = []
    heading, buf = "", []

    def flush():
        body = "\n".join(buf).strip()
        if body:
            sections.append(Section(text=body, heading=heading))

    for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
        if re.match(r"^#{1,4}\s+", line):
            flush()
            heading, buf = line.lstrip("# ").strip(), []
        else:
            buf.append(line)
    flush()
    return sections


def parse_pdf(path: Path) -> list[Section]:
    from pypdf import PdfReader

    out = []
    for i, page in enumerate(PdfReader(str(path)).pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            out.append(Section(text=text, page=i))
    return out  # scanned PDFs yield nothing: they need OCR (see DESIGN.md)


def parse_docx(path: Path) -> list[Section]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    sections: list[Section] = []
    heading, buf = "", []

    def flush():
        body = "\n\n".join(buf).strip()
        if body:
            sections.append(Section(text=body, heading=heading))

    for child in doc.element.body.iterchildren():  # document order, so tables stay next to their heading
        if child.tag.endswith("}p"):
            p = Paragraph(child, doc)
            if not p.text.strip():
                continue
            if p.style is not None and p.style.name.lower().startswith(("heading", "title")):
                flush()
                heading, buf = p.text.strip(), []
            else:
                buf.append(p.text.strip())
        elif child.tag.endswith("}tbl"):
            rows = [" | ".join(c.text.strip() for c in r.cells) for r in Table(child, doc).rows]
            flush()
            buf = []
            sections.append(Section(text="\n".join(rows), heading=heading, is_table=True))
    flush()
    return sections


_COLUMN_META = {"batch_id": "batch_id", "model": "model", "supplier": "supplier", "date": "doc_date", "ticket_id": "ticket_id"}


def parse_xlsx(path: Path) -> list[Section]:
    """One Section per row ("col: value; ..."). Row-oriented numeric data belongs in SQL, not here."""
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            continue
        names = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(header)]
        for r in rows:
            if all(v is None for v in r):
                continue
            pairs = {n: ("" if v is None else str(v)[:10] if hasattr(v, "isoformat") else str(v)) for n, v in zip(names, r)}
            meta = {_COLUMN_META[k.lower()]: v for k, v in pairs.items() if k.lower() in _COLUMN_META and v}
            text = "; ".join(f"{k}: {v}" for k, v in pairs.items() if v)
            out.append(Section(text=text, heading=ws.title, meta=meta))
    return out


PARSERS = {".md": parse_text, ".txt": parse_text, ".pdf": parse_pdf, ".docx": parse_docx, ".xlsx": parse_xlsx}


def parse(path: Path) -> list[Section]:
    return PARSERS[path.suffix.lower()](path)
